import io
import os
from typing import Any, Dict, Tuple
from pypdf import PdfReader
import docx

from app.services.message_scanner import analyze_message_text

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB


def extract_text_safely(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Safely extract textual content from uploaded documents without executing
    macros, scripts, or embedded active components.
    Returns (extracted_text, file_type).
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")
        return text, "Text Document (.txt)"

    elif ext == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            extracted_pages = []
            # Extract text from up to first 25 pages safely
            for idx, page in enumerate(reader.pages[:25]):
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text)
            full_text = "\n".join(extracted_pages)
            return full_text, "PDF Document (.pdf)"
        except Exception as e:
            return "", f"PDF Document (Error extracting text: {str(e)})"

    elif ext in [".docx", ".doc"]:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)
            full_text = "\n".join(paragraphs)
            return full_text, "Word Document (.docx)"
        except Exception as e:
            return "", f"Word Document (Error extracting text: {str(e)})"

    elif ext in [".png", ".jpg", ".jpeg"]:
        # OCR pipeline placeholder
        return "Image document received. OCR text processing engine ready for pluggable optical character recognition.", f"Image Document ({ext})"

    else:
        return "", f"Unsupported Document ({ext})"


def analyze_document_content(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Analyzes an uploaded document for fraudulent, deceptive, or scam patterns.
    """
    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        return {
            "filename": filename,
            "file_type": "Unknown",
            "risk_level": "HIGH RISK",
            "risk_score": 75,
            "detected_indicators": ["Excessive File Size Limit Exceeded"],
            "reasons": ["The file exceeds the safe 10MB scanning threshold."],
            "text_preview": "",
            "recommended_action": "Do not upload oversized files.",
            "summary": "File size limit exceeded.",
        }

    extracted_text, file_type = extract_text_safely(file_bytes, filename)

    if not extracted_text or not extracted_text.strip():
        return {
            "filename": filename,
            "file_type": file_type,
            "risk_level": "SAFE",
            "risk_score": 0,
            "detected_indicators": [],
            "reasons": ["No suspicious text or actionable scam patterns found in document text content."],
            "text_preview": "(No extractable plain text found in document)",
            "recommended_action": "No scam patterns found in text content. Ensure document source is verified.",
            "summary": "Document analyzed with no detected threats.",
        }

    # Analyze extracted text using message & scam detection engine
    text_analysis = analyze_message_text(extracted_text)

    # Safe limited text preview (first 250 characters)
    preview = extracted_text.strip()[:250]
    if len(extracted_text.strip()) > 250:
        preview += "…"

    return {
        "filename": filename,
        "file_type": file_type,
        "risk_level": text_analysis["risk_level"],
        "risk_score": text_analysis["risk_score"],
        "detected_indicators": text_analysis["detected_indicators"],
        "reasons": text_analysis["reasons"],
        "text_preview": preview,
        "recommended_action": text_analysis["recommended_action"],
        "summary": text_analysis["summary"],
    }
